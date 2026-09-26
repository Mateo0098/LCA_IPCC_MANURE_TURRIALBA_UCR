from __future__ import annotations

import re
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Rectangle
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import generate_conclusions_docx as conclusions_source  # noqa: E402
import generate_methodology_docx as methodology_source  # noqa: E402
import generate_results_docx as results_source  # noqa: E402
from academic_text_utils import clean_academic_label  # noqa: E402
from master_word_format import (  # noqa: E402
    add_master_caption,
    apply_master_format,
    finalize_document_format,
    format_table_like_master,
)
from reference_docx_utils import (  # noqa: E402
    REGISTERED_REFERENCE_SHA256,
    assert_reference_docx_intact,
    get_reference_docx_path,
    sha256_file,
)


PROVISIONAL_LABEL = "PROVISIONAL M1–M2"
MASTER = ROOT / "MASTER_escrito" / "TFG_ACV_Estiercol_MASTER.docx"
REFERENCES = ROOT / "docs" / "REFERENCIAS_TFG.md"
SENSITIVITY = ROOT / "outputs" / "auditoria_n_reactivo" / "sensibilidad_n2_a2.csv"
OUT_DIR = ROOT / "outputs" / "documentos_tfg"
OUT_DOCX = OUT_DIR / "TFG_ACV_Estiercol_INTEGRAL_PROVISIONAL_M1_M2.docx"
OUT_VALIDATION = OUT_DIR / "reporte_validacion_documento_integral.md"
FIG_DIR = ROOT / "outputs" / "graficos_tesis"
INTEGRAL_FIG_DIR = OUT_DIR / "recursos_integral"
SYSTEM_BOUNDARY_PNG = INTEGRAL_FIG_DIR / "fronteras_sistema_integral.png"
SYSTEM_BOUNDARY_SVG = INTEGRAL_FIG_DIR / "fronteras_sistema_integral.svg"

OBJECTIVE_GENERAL = (
    "Desarrollar un Análisis de Ciclo de Vida del manejo del estiércol del ganado "
    "de una lechería especializada, para estimación del impacto ambiental."
)
OBJECTIVE_SPECIFIC_1 = (
    "Realizar el inventario para el Análisis de Ciclo de Vida del estiércol bovino "
    "manejado mediante lombricompostaje y el aplicado como purines directamente en "
    "los campos de pastoreo."
)
OBJECTIVE_SPECIFIC_2 = (
    "Evaluar el impacto ambiental del estiércol bovino manejado mediante "
    "lombricompostaje y el aplicado como purines directamente en los campos de pastoreo."
)

EXPECTED_HEADINGS = [
    "1. Introducción",
    "2. Marco teórico y antecedentes",
    "3. Objetivos",
    "4. Metodología",
    "5. Resultados provisionales M1–M2",
    "6. Discusión provisional",
    "7. Conclusiones provisionales",
    "8. Recomendaciones",
    "9. Referencias",
    "Apéndice A. Trazabilidad entre objetivos y evidencia provisional",
]

REQUIRED_REFERENCE_KEYS = {
    "Arfelli2023",
    "Baek2020",
    "Barrantes2019",
    "CalderonChaves2020",
    "Callejo2020",
    "Curran2006",
    "Curran2013",
    "EMEPEEA2023",
    "EuropeanCommission2022",
    "Fernandez2014",
    "Garro2016",
    "Garza2021",
    "Hou2015",
    "IMN2021",
    "IMN2026",
    "INEC2023",
    "IPCC2019",
    "ISO140402006",
    "Jjagwe2019",
    "Komakech2014",
    "Komakech2015",
    "Komakech2016",
    "Kupper2020",
    "Lim2016",
    "Macktoobian2024",
    "MAG2018",
    "MARM2010",
    "Moller2004",
    "NRCS2009",
    "RojasElizondo2020",
    "Salazar2012",
    "SanchezBrenes2026",
    "VargasSarmiento2023",
    "VanderZaag2013",
    "VanderZaag2018",
    "Zhou2017",
}


@dataclass
class EditorialCounters:
    table: int = 0
    figure: int = 0
    equation: int = 0


def validate_inputs() -> Path:
    reference = get_reference_docx_path(ROOT)
    if reference != MASTER:
        raise RuntimeError("La ruta validada del MASTER no coincide con la esperada.")
    required = [REFERENCES, SENSITIVITY]
    missing = [path.relative_to(ROOT) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError(f"Faltan entradas para el documento integral: {missing}")
    methodology_source.validate_inputs()
    results_source.validate_inputs()
    return reference


def master_text(document: Document, index: int) -> str:
    value = document.paragraphs[index].text.strip()
    if not value:
        raise RuntimeError(f"El párrafo {index} esperado del MASTER está vacío.")
    return value


def add_text(document: Document, paragraphs: list[str]) -> None:
    for value in paragraphs:
        document.add_paragraph(clean_academic_label(value), style="Normal")


def add_master_paragraphs(document: Document, master: Document, indexes: range | list[int]) -> None:
    values = [master.paragraphs[index].text.strip() for index in indexes]
    add_text(document, [value for value in values if value])


def add_chapter(document: Document, title: str, *, first: bool = False) -> None:
    if not first:
        document.add_page_break()
    document.add_heading(title, level=1)


def display_value(value: object, decimals: int = 4) -> str:
    if pd.isna(value):
        return "—"
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, (int, float)):
        return results_source.fmt(value, decimals)
    text = clean_academic_label(value)
    replacements = {
        "ESTIERCOL FRESCO": "Estiércol fresco",
        "SOL: PRECOMPOSTADO": "Estiércol precompostado",
        "LIQ: AGUA VERDE": "Aguas verdes",
        "LIQ: PURINES": "Purines",
        "ipcc": "IPCC",
        "Amoniaco": "Amoníaco",
        "NOx as NO₂": "NOx como NO₂",
        "15a edición": "15.ª edición",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text


def add_dataframe(
    document: Document,
    profile,
    counters: EditorialCounters,
    description: str,
    dataframe: pd.DataFrame,
    *,
    decimals: int = 4,
) -> int:
    counters.table += 1
    caption_start = len(document.paragraphs)
    add_master_caption(document, f"Tabla {counters.table}. {description}")
    for paragraph in document.paragraphs[caption_start:]:
        paragraph.paragraph_format.keep_with_next = True
    table = document.add_table(rows=1, cols=len(dataframe.columns))
    for column_index, column in enumerate(dataframe.columns):
        table.rows[0].cells[column_index].text = clean_academic_label(column)
    for _, row in dataframe.iterrows():
        cells = table.add_row().cells
        for column_index, value in enumerate(row):
            cells[column_index].text = display_value(value, decimals)
    for row_index, row in enumerate(table.rows):
        tr_properties = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_properties.append(cant_split)
        if row_index == 0:
            table_header = OxmlElement("w:tblHeader")
            table_header.set(qn("w:val"), "true")
            tr_properties.append(table_header)
    format_table_like_master(table, profile)
    return counters.table


def add_figure(
    document: Document,
    counters: EditorialCounters,
    file_name: str | Path,
    description: str,
) -> int:
    image = Path(file_name)
    if not image.is_absolute():
        image = FIG_DIR / image
    if not image.exists():
        raise FileNotFoundError(f"No existe la figura requerida: {image.relative_to(ROOT)}")
    counters.figure += 1
    caption_start = len(document.paragraphs)
    add_master_caption(document, f"Figura {counters.figure}. {description}")
    for caption_paragraph in document.paragraphs[caption_start:]:
        caption_paragraph.paragraph_format.keep_with_next = True
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image), width=Inches(6.0))
    return counters.figure


def add_equation(
    document: Document,
    counters: EditorialCounters,
    equation: str,
    definition: str | None = None,
) -> int:
    counters.equation += 1
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"{equation}    ({counters.equation})")
    run.font.name = "Cambria Math"
    if definition:
        document.add_paragraph(definition, style="Normal")
    return counters.equation


def generate_system_boundary_figure() -> None:
    """Reconstruye las fronteras vigentes de ambos escenarios desde código."""

    def box(axis, x, y, width, height, text, *, facecolor="#EAF2F8", fontsize=8.5):
        patch = FancyBboxPatch(
            (x, y),
            width,
            height,
            boxstyle="round,pad=0.015,rounding_size=0.012",
            linewidth=1.0,
            edgecolor="#1F3A4D",
            facecolor=facecolor,
        )
        axis.add_patch(patch)
        axis.text(x + width / 2, y + height / 2, text, ha="center", va="center", fontsize=fontsize)
        return patch

    def arrow(axis, start, end, *, text=None, text_offset=(0.0, 0.0), style="-"):
        axis.add_patch(
            FancyArrowPatch(
                start,
                end,
                arrowstyle="-|>",
                mutation_scale=12,
                linewidth=1.2,
                linestyle=style,
                color="#263238",
                shrinkA=2,
                shrinkB=2,
            )
        )
        if text:
            axis.text(
                (start[0] + end[0]) / 2 + text_offset[0],
                (start[1] + end[1]) / 2 + text_offset[1],
                text,
                ha="center",
                va="center",
                fontsize=7.5,
                color="#263238",
            )

    def prepare_axis(axis, title):
        axis.set_xlim(0, 1)
        axis.set_ylim(0, 1)
        axis.axis("off")
        axis.add_patch(Rectangle((0.025, 0.07), 0.95, 0.84, fill=False, linewidth=1.4, edgecolor="#2E75B6"))
        axis.text(0.5, 0.955, title, ha="center", va="top", fontsize=11, fontweight="bold")
        axis.text(0.965, 0.925, "Frontera del sistema", ha="right", va="bottom", fontsize=7.5, color="#2E75B6")

    fig, axes = plt.subplots(2, 1, figsize=(10, 9.2))

    ax = axes[0]
    prepare_axis(ax, "Escenario A: ruta sólida y ruta de aguas verdes")
    box(ax, 0.36, 0.76, 0.28, 0.09, "Flujo común de estiércol fresco\n26 278,7 kg/año", facecolor="#F3F6F8")
    ax.text(0.50, 0.69, "Separación física durante la limpieza", ha="center", va="center", fontsize=7.5, fontstyle="italic")
    box(ax, 0.11, 0.49, 0.25, 0.10, "A1: Precomposteo")
    box(ax, 0.11, 0.21, 0.25, 0.10, "A2: Lombricompostaje")
    box(ax, 0.64, 0.49, 0.25, 0.10, "A3: Almacenamiento\nde aguas verdes")
    box(ax, 0.64, 0.21, 0.25, 0.10, "A4: Aplicación de aguas verdes\nen campos de pastoreo")
    arrow(ax, (0.46, 0.76), (0.235, 0.59), text="Fracción paleada\n17 525,1 kg/año", text_offset=(-0.06, 0.035))
    arrow(ax, (0.54, 0.76), (0.765, 0.59), text="Fracción remanente\n8 753,6 kg/año", text_offset=(0.06, 0.035))
    arrow(ax, (0.235, 0.49), (0.235, 0.31))
    arrow(ax, (0.765, 0.49), (0.765, 0.31))
    arrow(ax, (0.97, 0.54), (0.89, 0.54), text="Agua de lavado", text_offset=(0.0, 0.045))
    arrow(ax, (0.11, 0.525), (0.055, 0.40), text="Drenaje a suelo agrícola\no matorral", text_offset=(0.035, -0.01), style="--")
    ax.text(0.90, 0.64, "Electricidad", ha="center", va="center", fontsize=7.5)
    arrow(ax, (0.88, 0.62), (0.85, 0.59))
    ax.text(0.765, 0.14, "Diésel", ha="center", va="center", fontsize=7.5)
    arrow(ax, (0.765, 0.16), (0.765, 0.21))
    ax.text(0.235, 0.13, "Lombricompost", ha="center", va="center", fontsize=7.5)
    arrow(ax, (0.235, 0.21), (0.235, 0.16))
    ax.text(0.50, 0.085, "Las emisiones de manejo se cuantifican en cada etapa.", ha="center", va="center", fontsize=7.5)

    ax = axes[1]
    prepare_axis(ax, "Escenario B: almacenamiento y aplicación de purines")
    box(ax, 0.08, 0.54, 0.25, 0.12, "Flujo común de estiércol fresco\n26 278,7 kg/año", facecolor="#F3F6F8")
    box(ax, 0.40, 0.54, 0.22, 0.12, "B1: Almacenamiento\nde purines")
    box(ax, 0.70, 0.54, 0.24, 0.12, "B2: Aplicación de purines\nen campo de pastoreo")
    arrow(ax, (0.33, 0.60), (0.40, 0.60))
    arrow(ax, (0.62, 0.60), (0.70, 0.60))
    arrow(ax, (0.51, 0.80), (0.51, 0.66), text="Agua de lavado", text_offset=(0.09, 0.0))
    ax.text(0.46, 0.43, "Electricidad", ha="center", va="center", fontsize=7.5)
    arrow(ax, (0.46, 0.46), (0.46, 0.54))
    ax.text(0.77, 0.43, "Diésel", ha="center", va="center", fontsize=7.5)
    arrow(ax, (0.77, 0.46), (0.77, 0.54))
    ax.text(0.50, 0.22, "Se cuantifican emisiones de manejo y consumos operativos dentro de la frontera.", ha="center", va="center", fontsize=8)
    arrow(ax, (0.57, 0.54), (0.57, 0.31), text="Emisiones", text_offset=(0.065, 0.0), style="--")
    arrow(ax, (0.88, 0.54), (0.88, 0.31), text="Emisiones", text_offset=(0.065, 0.0), style="--")

    fig.subplots_adjust(left=0.03, right=0.97, top=0.98, bottom=0.03, hspace=0.10)
    INTEGRAL_FIG_DIR.mkdir(parents=True, exist_ok=True)
    fig.savefig(SYSTEM_BOUNDARY_PNG, dpi=300, bbox_inches="tight", metadata={"Date": None})
    fig.savefig(SYSTEM_BOUNDARY_SVG, bbox_inches="tight", metadata={"Date": None})
    plt.close(fig)
    svg_lines = SYSTEM_BOUNDARY_SVG.read_text(encoding="utf-8").splitlines()
    SYSTEM_BOUNDARY_SVG.write_text(
        "\n".join(line.rstrip() for line in svg_lines) + "\n",
        encoding="utf-8",
    )


def read_reference_registry() -> list[dict[str, str]]:
    entries: list[dict[str, str]] = []
    for line in REFERENCES.read_text(encoding="utf-8").splitlines():
        if not line.startswith("|"):
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if len(cells) != 6 or cells[0] in {"Clave", "---"} or set(cells[0]) == {"-"}:
            continue
        entries.append(
            {
                "key": cells[0],
                "reference": cells[1].replace("*", "").replace("`", ""),
                "section": cells[2],
                "purpose": cells[3],
                "status": cells[4],
                "notes": cells[5],
            }
        )
    found = {entry["key"] for entry in entries}
    missing = sorted(REQUIRED_REFERENCE_KEYS - found)
    if missing:
        raise RuntimeError(f"Faltan referencias requeridas en el registro: {missing}")
    return entries


def conclusion_items() -> list[dict[str, str]]:
    totals = conclusions_source.impact_totals()
    percentages = conclusions_source.comparisons(totals)
    reference_flow, normalized = conclusions_source.processed_indicators(totals)
    collected, remainder = conclusions_source.flow_inventory(reference_flow)
    stages = conclusions_source.stage_totals()
    conclusions_source.validate_stage_sums(stages, totals)
    return conclusions_source.build_conclusions(
        totals,
        normalized,
        percentages,
        stages,
        reference_flow,
        collected,
        remainder,
    )


def sensitivity_summary() -> tuple[pd.DataFrame, float]:
    source = pd.read_csv(SENSITIVITY, encoding="utf-8-sig")
    table = pd.DataFrame(
        {
            "Factor de N₂-N respecto al TAN": source["emep_n2_factor_fraction_tan"],
            "N₂-N estimado (kg)": source["n2_n_kg"],
            "N total remanente (kg)": source["n_total_out_kg"],
            "Cambio respecto al caso central (% del N de entrada)": source[
                "change_vs_central_pct_n_total_in"
            ],
        }
    )
    maximum = float(source["change_vs_central_pct_n_total_in"].abs().max())
    return table, maximum


def add_title_page(document: Document, profile) -> None:
    centered = document.styles.add_style("Portada centrada", WD_STYLE_TYPE.PARAGRAPH)
    centered.base_style = document.styles["Normal"]
    centered.font.name = profile.font_name
    centered.font.size = Pt(profile.body_size_pt)
    centered.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    centered.paragraph_format.left_indent = Pt(0)
    centered.paragraph_format.first_line_indent = Pt(0)

    title_style = document.styles.add_style("Título integral", WD_STYLE_TYPE.PARAGRAPH)
    title_style.base_style = centered
    title_style.font.name = profile.font_name
    title_style.font.size = Pt(profile.title_size_pt)
    title_style.font.italic = True
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_lines = [
        "Universidad de Costa Rica",
        "Escuela de Ingeniería de Biosistemas",
        "Trabajo Final de Graduación",
    ]
    for line in title_lines:
        paragraph = document.add_paragraph(style="Portada centrada")
        paragraph.add_run(line).bold = True
    document.add_paragraph("")
    document.add_paragraph(
        "Análisis de ciclo de vida de los desechos bovinos, sólidos y líquidos "
        "producidos en una lechería especializada en Turrialba, Costa Rica",
        style="Título integral",
    )
    document.add_paragraph("")
    author = document.add_paragraph(style="Portada centrada")
    author.add_run("Mateo Cerdas Barboza\nCarné B71946")
    document.add_paragraph("")
    status = document.add_paragraph(style="Portada centrada")
    run = status.add_run(PROVISIONAL_LABEL)
    run.bold = True
    note = document.add_paragraph(style="Portada centrada")
    note.add_run(
        "Documento integral de trabajo para revisión académica. La jornada M3 permanece pendiente."
    ).italic = True
    date = document.add_paragraph(style="Portada centrada")
    date.add_run("Turrialba, Costa Rica\nSeptiembre de 2026")


def add_preliminaries(document: Document) -> None:
    document.add_page_break()
    document.add_heading("Estado del documento", level=1)
    add_text(
        document,
        [
            "Este documento integra la propuesta académica original con la metodología ejecutada y los productos regenerables de la corrida vigente. Los resultados, la discusión y las conclusiones corresponden a la integración PROVISIONAL M1–M2. La incorporación de M3 actualizará la caracterización y puede modificar las magnitudes, comparaciones e interpretaciones presentadas.",
            "El documento maestro aprobado se conserva como fuente protegida de continuidad académica y formato. La presente versión establece una ruta editorial única hacia el TFG final sin alterar ese documento original.",
        ],
    )
    document.add_heading("Contenido", level=1)
    for heading in EXPECTED_HEADINGS:
        document.add_paragraph(heading, style="Normal")


def build_document() -> tuple[int, int, int, int]:
    master = Document(str(MASTER))
    document = Document()
    profile = apply_master_format(document, MASTER)
    counters = EditorialCounters()
    methodology_context = methodology_source.methodology_context()
    result_context = results_source.results_context()
    references = read_reference_registry()

    for section in document.sections:
        header = section.header.paragraphs[0]
        header.text = PROVISIONAL_LABEL
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_title_page(document, profile)
    add_preliminaries(document)

    add_chapter(document, "1. Introducción")
    document.add_heading("1.1 Justificación", level=2)
    add_master_paragraphs(document, master, range(28, 37))
    add_text(
        document,
        [
            "Las cifras contextuales y referencias heredadas de la propuesta se conservan en esta versión para mantener continuidad académica; su cotejo bibliográfico final se encuentra identificado en el registro integral de referencias.",
        ],
    )
    document.add_heading("1.2 Delimitación del problema", level=2)
    add_master_paragraphs(document, master, [39, 40])
    add_text(
        document,
        [
            "El estudio aplicó el ACV a dos alternativas del manejo de estiércol en la lechería: el Escenario A, que integra la ruta sólida de precomposteo y lombricompostaje con la ruta de aguas verdes, y el Escenario B, que representa el almacenamiento y la aplicación directa de purines. La evaluación abarcó cambio climático, eutrofización terrestre y eutrofización marina. En esta etapa documental, la evidencia disponible integra M1 y M2 y mantiene pendiente M3.",
        ],
    )

    add_chapter(document, "2. Marco teórico y antecedentes")
    theoretical_blocks = [
        ("2.1 Generación de excreta bovina", range(50, 55)),
        ("2.2 Métodos para el manejo de la excreta bovina", []),
        ("2.2.1 Lombricompostaje", [58]),
        ("2.2.2 Enmienda agrícola", [61]),
        ("2.3 Análisis de ciclo de vida e impacto ambiental", range(64, 72)),
        ("2.4 Análisis de ciclo de vida del manejo del estiércol bovino", range(74, 78)),
        ("2.5 Antecedentes", range(79, 88)),
    ]
    for title, indexes in theoretical_blocks:
        document.add_heading(title, level=2 if title.count(".") == 1 else 3)
        if indexes:
            add_master_paragraphs(document, master, indexes)

    add_chapter(document, "3. Objetivos")
    document.add_heading("3.1 Objetivo general", level=2)
    document.add_paragraph(master_text(master, 91), style="Normal")
    document.add_heading("3.2 Objetivos específicos", level=2)
    document.add_paragraph(master_text(master, 93), style="Normal")
    document.add_paragraph(master_text(master, 94), style="Normal")

    add_chapter(document, "4. Metodología")
    document.add_heading("4.1 Enfoque metodológico y sitio de estudio", level=2)
    add_text(
        document,
        [
            "El estudio se desarrolló como un Análisis de Ciclo de Vida aplicado al manejo del estiércol bovino en una lechería especializada de Turrialba, Costa Rica. Comprendió la definición de meta y alcance, la construcción del Inventario de Ciclo de Vida, la estimación de emisiones por etapa y la conversión de esas emisiones a indicadores de impacto ambiental.",
        ],
    )
    add_master_paragraphs(document, master, range(100, 103))
    add_text(
        document,
        [
            "La descripción del sitio y de sus operaciones conserva la formulación de la propuesta cuando continúa siendo compatible con el trabajo ejecutado. Los parámetros operativos cuantitativos se actualizaron con la evidencia de la misma lechería reportada por Sánchez-Romero y Brenes-Gamboa (2026).",
            "En la operación habitual del Escenario A, el estiércol recogido con pala se conduce a A1: Precomposteo y después a A2: Lombricompostaje. El remanente del piso se incorpora al agua de lavado y las aguas verdes llegan mediante el drenaje o canal a la tanqueta para su almacenamiento en A3. La tanqueta se vacía aproximadamente cada tres días y el tractor acciona el cañón VAIA para la aplicación en A4. La tanqueta almacena; el cañón aplica.",
        ],
    )
    document.add_heading("4.2 Meta, alcance, unidad funcional y escenarios", level=2)
    add_text(
        document,
        [
            "La meta fue comparar el desempeño ambiental de dos alternativas de manejo bajo una misma unidad funcional de 1 kg de estiércol fresco manejado. Los flujos y emisiones anuales describen la escala operacional; los indicadores por kilogramo corresponden a la normalización respecto a la unidad funcional.",
            f"El flujo anual común fue {results_source.fmt(methodology_context['flujo_referencia'], 6)} kg de estiércol fresco/año. En el Escenario A, la fracción sólida ingresó a A1: Precomposteo y continuó hacia A2: Lombricompostaje; el remanente se incorporó a A3: Almacenamiento de aguas verdes y A4: Aplicación de aguas verdes en campos de pastoreo. En el Escenario B, el flujo completo ingresó a B1: Almacenamiento de purines y continuó hacia B2: Aplicación de purines en campo de pastoreo.",
            "El Escenario B no fue la operación habitual permanente ni una alternativa puramente hipotética. Para materializarlo temporalmente se suspendió la desviación normal del sólido hacia precomposteo y lombricompostaje, el estiércol paleado se dirigió a la tanqueta y el remanente del piso se incorporó mediante lavado. Los purines resultantes fueron acumulados, observados y muestreados físicamente. Las campañas A y B se realizaron en momentos distintos con la misma tanqueta y no implicaron la coexistencia de ambos contenidos. B1 representa el almacenamiento y B2 la aplicación con el mismo conjunto tractor–cañón utilizado para A4.",
            "A1 duró aproximadamente entre 21 días y cerca de un mes, por lo que se representa como tres a cuatro semanas sin tratar 28 días como una medición exacta. A2 comenzó después de A1, cuando el material precompostado ingresó a las camas, y su operación regular duró aproximadamente 13 semanas. Vargas Sarmiento (2023, sección 5.2.2.1, p. 14; sección 6.1.3, p. 25) documentó en el mismo lombricario 13 semanas desde la siembra de las lombrices y para procesar toda la boñiga. Estas duraciones son contextuales y no escalan factores. El presente TFG no muestreó lombricompost terminado; las muestras correspondieron a material precompostado previo a A2.",
        ],
    )
    generate_system_boundary_figure()
    boundary_figure = add_figure(
        document,
        counters,
        SYSTEM_BOUNDARY_PNG,
        "Fronteras y conexiones físicas de los escenarios evaluados.",
    )
    add_text(
        document,
        [
            f"La Figura {boundary_figure} reconstruye los diagramas conceptuales de la propuesta con la nomenclatura y las conexiones vigentes. En el Escenario A, la fracción paleada y la fracción remanente siguen rutas físicamente separadas; en el Escenario B, el flujo común completo continúa por almacenamiento y aplicación. Los consumos de electricidad y diésel se asignan a las etapas operativas correspondientes.",
        ],
    )
    stage_table = methodology_source.stage_summary().drop(
        columns=["Modelo de estimación"], errors="ignore"
    )
    table_1 = add_dataframe(
        document,
        profile,
        counters,
        "Etapas oficiales de los escenarios evaluados.",
        stage_table,
        decimals=2,
    )
    add_text(document, [f"La Tabla {table_1} presenta la jerarquía de etapas utilizada en todo el documento integral."])

    document.add_heading("4.3 Muestreo e integración temporal", level=2)
    add_text(
        document,
        [
            "La jerarquía estadística fue réplica analítica, muestra compuesta, promedio de jornada e integración entre jornadas. Las réplicas analíticas no se trataron como observaciones temporales independientes y las jornadas recibieron igual peso temporal.",
            "En M1 se analizaron, por cada sólido, dos muestras compuestas en Bioenergía y otras dos muestras compuestas físicamente independientes en el laboratorio externo: LASA para estiércol fresco y CIA para precompostado. En M2 se conservaron tres muestras compuestas por sólido; Bioenergía realizó tres réplicas gravimétricas por muestra y el remanente de esas mismas muestras fue analizado por LASA o CIA.",
            "Bioenergía determinó humedad y materia seca por gravimetría a 105 °C durante 16 h. El CIA determinó N y C del precompostado por Dumas sobre muestra seca o acondicionada a 80 °C durante 48 h. El porcentaje de N se combinó con la materia seca independiente de Bioenergía para construir el benchmark húmedo de A2; C y C/N permanecieron como caracterización descriptiva, sin conversión húmeda ni uso productivo.",
            "Para los sólidos metodológicamente comparables, la integración provisional combinó M1 y M2. En aguas verdes y purines, M1 correspondió a especiación y se conservó para trazabilidad; el N total líquido activo procedió de M2 mediante Kjeldahl. M3 permanece pendiente y se incorporará mediante el mismo pipeline para producir la caracterización final.",
            "La transformación de estiércol fresco a precompostado se calculó primero por jornada mediante materia seca y cenizas de ambos materiales; posteriormente se integraron los factores de jornada con igual peso temporal. La pérdida integrada se derivó del factor integrado.",
            "Antes de cada campaña líquida, la tanqueta se vació el jueves por la tarde y acumuló material hasta el lunes por la mañana. Durante ese intervalo de acumulación previo al muestreo continuaron las entradas y ocurrieron dos lavados; no fue una carga cerrada ni un ensayo de almacenamiento estático. Esta preparación de muestra no sustituyó la frecuencia operativa representativa de aproximadamente tres días.",
        ],
    )
    characterization = results_source.characterization_summary()
    table_2 = add_dataframe(
        document,
        profile,
        counters,
        "Caracterización fisicoquímica provisional de los materiales analizados.",
        characterization,
        decimals=3,
    )
    add_text(document, [f"La Tabla {table_2} resume los datos experimentales promovidos a la corrida provisional."])

    document.add_heading("4.4 Balance secuencial de nitrógeno y estimación de emisiones", level=2)
    add_text(
        document,
        [
            "El N total constituyó el balance físico principal y el nitrógeno amoniacal total un subbalance sujeto a 0 ≤ TAN ≤ N total. TAN se inicializó como 0,60 del N total únicamente en las fronteras de estiércol fresco y ambos componentes se propagaron entre etapas físicamente conectadas.",
            "En particular, A2 recibió el N total y el TAN remanentes de A1. La medición de N del precompostado se mantuvo como benchmark experimental y no reinicializó estos flujos productivos.",
            "Las pérdidas explícitas de NH₃-N, NOx-N y N₂-N definidas por EMEP/EEA redujeron el TAN y el N total. El N₂O-N directo y las pérdidas hídricas definidas por IPCC redujeron el N total una sola vez. El N₂O indirecto por volatilización se calculó con las especies explícitas NH₃-N y NOx-N; el NO₃⁻ se originó únicamente en rutas hídricas justificadas.",
            "El ledger fue secuencial: A2 recibió el N total y el pool residual de TAN a la salida de A1. Por ello, aplicar factores en ambas etapas no duplicó matemáticamente los pools originales. FracGasMS permaneció exclusivamente como benchmark y no generó otra volatilización ni alimentó EF4. Las duraciones específicas introducen incertidumbre de representatividad temporal y de transferencia de los proxies, pero no justifican escalado lineal de los factores.",
            "A2: Lombricompostaje se representó mediante una arquitectura híbrida explícita y trazable: la categoría IPCC de compostaje en hileras pasivas se utilizó como proxy para CH₄ y N₂O directo; Komakech et al. (2016), como proxy experimental aprobado para NH₃; y EMEP/EEA de almacenamiento sólido, como proxy metodológico aprobado provisionalmente para NO y N₂. La fracción de pérdida de N por lixiviación se estableció en cero para las condiciones operativas modeladas, sin cambiar el factor genérico de la categoría ni afirmar imposibilidad física universal. Las ecuaciones siguientes documentan el núcleo necesario para reproducir la lógica vigente.",
        ],
    )
    document.add_heading("4.4.1 Metano de las etapas de manejo", level=3)
    add_equation(
        document,
        counters,
        "m_CH₄ = m_manejada × VS_húmeda × B₀ × 0,67 × (MCF/100) × AWMS",
        "m_CH₄ es la emisión de metano de la etapa; m_manejada es la masa húmeda manejada; VS_húmeda es la fracción de sólidos volátiles en base húmeda; B₀ es la capacidad máxima de producción de metano; MCF es el factor de conversión de metano y AWMS es la fracción manejada por el sistema.",
    )
    document.add_heading("4.4.2 Balance secuencial de N total y TAN", level=3)
    add_equation(
        document,
        counters,
        "TAN_fresco = 0,60 × N_total,fresco",
        "TAN_fresco es el nitrógeno amoniacal total en la frontera fresca y N_total,fresco es el nitrógeno total del estiércol fresco.",
    )
    add_equation(
        document,
        counters,
        "TAN_disponible = TAN_entrada + N_mineralizado",
        "La mineralización se aplica donde corresponde antes de estimar las pérdidas EMEP/EEA sobre el TAN disponible.",
    )
    add_equation(
        document,
        counters,
        "N_j = TAN_disponible × f_j,  j ∈ {NH₃-N, NO-N, N₂-N}",
        "La expresión resume las rutas parametrizadas con factores EMEP/EEA. En A2, la masa de NH₃ se estima una sola vez con el proxy experimental aprobado de Komakech et al. (2016), aplicado a la masa húmeda inferida de residuo orgánico que ingresa a la etapa, mientras NO-N y N₂-N conservan los proxies sólidos EMEP/EEA aprobados provisionalmente.",
    )
    add_equation(document, counters, "N_N₂O–N,directo = N_total,entrada × EF₃")
    add_equation(document, counters, "m_N₂O,directo = N_N₂O–N,directo × 44/28")
    add_equation(
        document,
        counters,
        "N_total,salida = N_total,entrada − N_NH₃ − N_NOx − N_N₂ − N_N₂O–N,directo − N_pérdida,hídrica",
        "Esta identidad expresa el cierre secuencial de N total en las etapas de manejo; cada pérdida física se descuenta una sola vez.",
    )
    add_equation(document, counters, "N_precursor,vol = N_NH₃ + N_NOx")
    add_equation(document, counters, "m_N₂O,ind,vol = N_precursor,vol × EF₄ × 44/28")
    document.add_heading("4.4.3 Rutas de N hacia el suelo y aplicación", level=3)
    add_text(
        document,
        [
            "A4 y B2 recibieron el N total y el TAN propagados desde sus etapas de almacenamiento. La volatilización de NH₃ se estimó sobre el TAN aplicado; el NOx, expresado como NO₂, se estimó sobre el N aplicado y se convirtió a masa de N. El N₂O directo y la lixiviación o escorrentía conservaron como base el N de estiércol aplicado.",
        ],
    )
    add_equation(
        document,
        counters,
        "N_drenaje,A1 = N_total,A1 × FracLeachMS_A1",
        "El drenaje de A1 se trata como entrada de N al suelo y no como una masa de NO₃⁻ directa.",
    )
    add_equation(
        document,
        counters,
        "N_NH₃,aplic = TAN_aplicado × f_NH₃;  N_NOx,aplic = (N_aplicado × f_NO₂) × 14/46",
    )
    add_equation(document, counters, "m_N₂O,directo,suelo = N_aplicado × EF₁ × 44/28")
    add_equation(document, counters, "N_lix,esc = N_entrada,suelo × FracLEACH_suelo")
    add_equation(document, counters, "m_NO₃⁻ = N_lix,esc × 62/14")
    add_equation(document, counters, "m_N₂O,ind,lix = N_lix,esc × EF₅ × 44/28")

    document.add_heading("4.5 Evaluación de impactos y consumos operativos", level=2)
    add_text(
        document,
        [
            "Environmental Footprint 3.1 se aplicó a las emisiones directas. El cambio climático se expresó en kg CO₂-eq, la eutrofización terrestre en mol N-eq y la eutrofización marina en kg N-eq; las categorías se interpretaron por separado.",
            "La electricidad de la bomba se evaluó con el factor agregado de consumo del IMN para 2025 como aproximación temporal del patrón observado en 2026. La combustión de diésel se representó mediante masas físicas de CO₂ fósil, CH₄ fósil y N₂O obtenidas con factores IMN y caracterizadas con EF 3.1. No se incorporaron cadenas completas de fondo.",
            "Los factores, la asignación de sistemas y los consumos operativos corresponden a las decisiones metodológicas vigentes. Esta integración documental no recalculó ni modificó el ACV.",
            "En A3/B1, el MCF de 38 % se mantuvo como proxy IPCC conservador de la categoría tabulada de un mes para clima tropical húmedo. No corresponde a una medición específica del MCF para la residencia operativa de tres días y no se escaló como 38 % × 3/30. La mineralización EMEP del 10 % tampoco se escaló por tres días. El tiempo físico de operación, el intervalo específico de preparación de las muestras y la duración tabulada del proxy metodológico se trataron como conceptos distintos.",
        ],
    )
    add_equation(
        document,
        counters,
        "I_c = Σᵢ (m_i × CF_i,c)",
        "I_c es el indicador de la categoría c, m_i es la masa del flujo elemental i y CF_i,c es su factor de caracterización en esa categoría.",
    )
    add_equation(document, counters, "CC_total = CC_manejo + CC_electricidad + CC_diésel")
    factor_table = methodology_source.characterization_factors()
    table_3 = add_dataframe(
        document,
        profile,
        counters,
        "Factores de emisión IMN y caracterización EF 3.1.",
        factor_table,
        decimals=4,
    )
    figure_1 = add_figure(
        document,
        counters,
        "fig_04_flujos_masa_equivalente_total.png",
        "Masa equivalente total por etapa y escenario.",
    )
    add_text(
        document,
        [
            f"La Tabla {table_3} documenta los factores visibles principales y la Figura {figure_1} muestra la escala operacional de los flujos.",
        ],
    )

    document.add_heading("4.6 Supuestos, consistencia y limitaciones", level=2)
    add_text(
        document,
        [
            "Los supuestos dominantes incluyen la equivalencia entre litro de agua y kilogramo equivalente, la extrapolación anual de las operaciones, la generación teórica de estiércol durante la permanencia en sala, la conservación de cenizas, la asignación de sistemas de manejo y sus factores, la relación TAN/N inicial y la representación del almacenamiento líquido mediante un MCF de 38 %.",
            "La consistencia se controló mediante balances de masa y nitrógeno, normalización común, trazabilidad entre integración experimental y parámetros activos, sumas por etapa y escenario, y comprobaciones de dirección, signo, unidad, dominancia y redondeo de las comparaciones narrativas.",
            "Las limitaciones principales son la representatividad temporal de M1–M2, la ausencia de una medición directa del MCF para aproximadamente tres días de residencia, la transferibilidad de Komakech, la representatividad de las categorías IPCC y del factor EMEP de N₂ en A2, la masa húmeda inferida de A2 y la ausencia de un balance cerrado de agua y sólidos. Son incertidumbres científicas de la arquitectura aprobada, no decisiones metodológicas abiertas ni impedimentos para el modelo pre-M3.",
        ],
    )

    add_chapter(document, "5. Resultados provisionales M1–M2")
    add_text(
        document,
        [
            "Todos los resultados de este capítulo corresponden a la corrida PROVISIONAL M1–M2. No constituyen una caracterización definitiva ni anticipan los resultados de M3.",
        ],
    )
    document.add_heading("5.1 Caracterización y flujos del inventario", level=2)
    characterization_index = characterization.set_index("Tipo de muestra")
    fresh = characterization_index.loc["Estiércol fresco"]
    precomposted = characterization_index.loc["Estiércol precompostado"]
    add_text(
        document,
        [
            f"El estiércol fresco presentó {results_source.fmt(fresh['Humedad (%)'], 2)} % de humedad y {results_source.fmt(fresh['Materia seca (%)'], 2)} % de materia seca. El material precompostado presentó {results_source.fmt(precomposted['Humedad (%)'], 2)} % de humedad y {results_source.fmt(precomposted['Materia seca (%)'], 2)} % de materia seca.",
            "La distribución de flujos mantuvo el mismo flujo anual de referencia para ambos escenarios. La ruta A separó la fracción sólida recolectada de la fracción incorporada a las aguas verdes; la ruta B condujo el flujo completo al almacenamiento y posterior aplicación de purines.",
        ],
    )
    figure_2 = add_figure(
        document,
        counters,
        "fig_01_caracterizacion_humedad_materia_seca.png",
        "Humedad y materia seca promedio por tipo de muestra.",
    )
    add_text(document, [f"La Figura {figure_2} presenta la caracterización gravimétrica provisional."])

    document.add_heading("5.2 Emisiones estimadas", level=2)
    emissions = results_source.emissions_summary()
    table_4 = add_dataframe(
        document,
        profile,
        counters,
        "Emisiones anuales del manejo por escenario y sustancia.",
        emissions,
        decimals=4,
    )
    figure_3 = add_figure(
        document,
        counters,
        "fig_06_emisiones_ch4.png",
        "Emisiones anuales de CH₄ por etapa y escenario.",
    )
    add_text(
        document,
        [
            f"La Tabla {table_4} resume las emisiones del manejo y la Figura {figure_3} muestra la distribución de CH₄. Las contribuciones operativas de electricidad y diésel se incorporaron en el indicador de cambio climático, manteniendo separada su trazabilidad.",
        ],
    )

    document.add_heading("5.3 Impactos por etapa y por escenario", level=2)
    stage_impacts = results_source.impact_stage_summary()
    if {"Escenario", "Etapa", "Nombre de etapa"} <= set(stage_impacts.columns):
        stage_codes = stage_impacts["Escenario"].astype(str) + stage_impacts["Etapa"].astype(int).astype(str)
        stage_names = stage_impacts["Nombre de etapa"].astype(str).str.replace(
            r"^Etapa\s+\d+:\s*", "", regex=True
        )
        stage_impacts.insert(1, "Etapa del sistema", stage_codes + ": " + stage_names)
        stage_impacts = stage_impacts.drop(columns=["Etapa", "Nombre de etapa"])
    table_5 = add_dataframe(
        document,
        profile,
        counters,
        "Impactos ambientales por etapa.",
        stage_impacts,
        decimals=6,
    )
    figure_4 = add_figure(
        document,
        counters,
        "fig_11_impactos_cambio_climatico_etapa.png",
        "Cambio climático total por etapa y escenario.",
    )
    figure_5 = add_figure(
        document,
        counters,
        "fig_12_impactos_eutrofizacion_terrestre_etapa.png",
        "Eutrofización terrestre EF 3.1 por etapa y escenario.",
    )
    add_text(
        document,
        [
            f"La Tabla {table_5} presenta los impactos por etapa. En el Escenario A, {result_context['cg_dominant_a_name']} aportó {results_source.fmt(result_context['cg_dominant_a_percentage'], 2)} % del cambio climático; en el Escenario B, {result_context['cg_dominant_b_name']} aportó {results_source.fmt(result_context['cg_dominant_b_percentage'], 2)} %. Las Figuras {figure_4} y {figure_5} muestran la distribución por etapa de cambio climático y eutrofización terrestre.",
        ],
    )
    document.add_page_break()
    document.add_heading("5.3.1 Resultados totales y comparación entre escenarios", level=3)
    totals = results_source.total_impact_summary()
    table_6 = add_dataframe(
        document,
        profile,
        counters,
        "Impactos ambientales anuales por escenario.",
        totals.iloc[:, :4],
        decimals=6,
    )
    normalized_totals = totals[[totals.columns[0], *totals.columns[4:]]].copy()
    normalized_totals = normalized_totals.rename(
        columns={
            totals.columns[4]: "Cambio climático (kg CO₂-eq/kg)",
            totals.columns[5]: "Eutrofización terrestre (mol N-eq/kg)",
            totals.columns[6]: "Eutrofización marina (kg N-eq/kg)",
        }
    )
    table_7 = add_dataframe(
        document,
        profile,
        counters,
        "Impactos ambientales normalizados por kilogramo de estiércol fresco manejado.",
        normalized_totals,
        decimals=6,
    )
    comparisons = results_source.comparison_summary()
    table_8 = add_dataframe(
        document,
        profile,
        counters,
        "Comparación de impactos ambientales entre escenarios.",
        comparisons,
        decimals=4,
    )
    figure_6 = add_figure(
        document,
        counters,
        "fig_17_comparacion_diferencia_porcentual.png",
        "Diferencia porcentual del Escenario B respecto al Escenario A por categoría de impacto.",
    )
    add_text(
        document,
        [
            f"La Tabla {table_6} presenta la magnitud anual y la Tabla {table_7} muestra los indicadores por unidad funcional. La Tabla {table_8} y la Figura {figure_6} presentan la comparación entre escenarios bajo la misma base funcional.",
        ],
    )

    document.add_heading("5.4 Contraste bibliográfico de A2", level=2)
    benchmark = results_source.a2_benchmark_summary()
    table_9 = add_dataframe(
        document,
        profile,
        counters,
        "Contraste bibliográfico de A2 sobre materia seca de entrada.",
        benchmark,
        decimals=6,
    )
    add_text(
        document,
        [
            f"La Tabla {table_9} contrasta las estimaciones oficiales de A2: Lombricompostaje con Jjagwe et al. (2019) sobre una base material armonizada. El contraste apoya la interpretación, pero no sustituye el inventario oficial ni constituye una validación formal del modelo.",
        ],
    )

    add_chapter(document, "6. Discusión provisional")
    document.add_heading("6.1 Interpretación de la comparación", level=2)
    add_text(
        document,
        [
            f"Bajo la unidad funcional común, el {result_context['cg_comparison'].higher_label} presentó el mayor cambio climático y el {result_context['cg_comparison'].lower_label} el menor. La diferencia B menos A fue {results_source.fmt(result_context['cg_percentage'], 2)} % respecto al Escenario A.",
            f"En eutrofización terrestre, el {result_context['et_comparison'].higher_label} presentó el mayor indicador; en eutrofización marina, el {result_context['eu_comparison'].higher_label} presentó el mayor indicador. Estas comparaciones corresponden a categorías distintas y no se agregan entre sí.",
            "La concentración del impacto en etapas distintas confirma que la interpretación debe considerar la estructura de cada alternativa, el almacenamiento, la aplicación al suelo y los consumos operativos. La evidencia provisional no permite generalizar una superioridad universal fuera de la lechería y de las condiciones modeladas.",
        ],
    )
    document.add_heading("6.2 Contraste con la literatura", level=2)
    add_text(
        document,
        [
            "El contraste con Jjagwe et al. (2019) mostró que las diferencias entre la estimación IPCC y la referencia experimental no siguieron una dirección uniforme para CH₄ y N₂O directo. Las diferencias de especie de lombriz, acondicionamiento, alimentación, humedad, duración, clima, escala y frecuencia de medición impiden interpretar esa comparación como equivalencia física entre sistemas.",
            "Las referencias de Komakech et al. (2016), Møller et al. (2004) y VanderZaag et al. (2013), junto con las directrices IPCC y EMEP/EEA, permiten trazar la selección y la interpretación de los factores empleados, pero no demuestran equivalencia física entre sus sistemas de origen y el sistema estudiado. En particular, el factor de NH₃ de Komakech procede de reactores pequeños en Kampala, con residuo predominantemente ganadero, concentraciones gaseosas medidas y un flujo de aire estimado mediante el índice respirométrico dinámico; su aplicación a la masa húmeda inferida de entrada a A2 es un proxy experimental aprobado y una extrapolación con diferencias de escala, alimentación, clima, duración y operación. Estas limitaciones no convierten su selección en una decisión abierta para la fase pre-M3.",
        ],
    )
    document.add_heading("6.3 Sensibilidad y consistencia", level=2)
    sensitivity, maximum_change = sensitivity_summary()
    table_10 = add_dataframe(
        document,
        profile,
        counters,
        "Comprobación existente de sensibilidad del factor de N₂ en A2.",
        sensitivity,
        decimals=4,
    )
    add_text(
        document,
        [
            f"La comprobación disponible para A2 varió el factor de N₂-N respecto al TAN entre 0 y 0,30. Frente al caso central de 0,30, el cambio máximo observado en el N total remanente equivalió a {results_source.fmt(maximum_change, 2)} % del N total de entrada. La Tabla {table_10} documenta esta prueba de QA; no constituye un análisis de sensibilidad completo de los impactos.",
            "Los supuestos con mayor capacidad material de afectar la comparación son el MCF del almacenamiento líquido, la caracterización de materia seca, sólidos volátiles y N después de M3, la relación TAN/N inicial, los factores de A2, la duración y frecuencia de los consumos operativos y los factores asociados a la aplicación al suelo.",
            "Después de M3 y antes de la interpretación final podrá evaluarse si una sensibilidad acotada de los proxies de alta influencia, en particular N₂ de A2 y NH₃ de Komakech, aporta valor suficiente. Su ausencia no impide validar el modelo pre-M3. La consistencia continuará verificándose mediante cierre de balances, igualdad del flujo funcional e integridad temporal de la corrida.",
        ],
    )

    add_chapter(document, "7. Conclusiones provisionales")
    add_text(
        document,
        [
            "Las siguientes conclusiones corresponden exclusivamente a la corrida PROVISIONAL M1–M2. M3 puede modificar las magnitudes, las etapas dominantes y la comparación; por tanto, estas conclusiones no representan el cierre experimental del TFG.",
        ],
    )
    conclusions = conclusion_items()
    for index, item in enumerate(conclusions, start=1):
        paragraph = document.add_paragraph(style="Normal")
        paragraph.add_run(f"Conclusión provisional {index}. ").bold = True
        paragraph.add_run(item["text"])

    add_chapter(document, "8. Recomendaciones")
    add_text(
        document,
        [
            "Incorporar M3 mediante la misma secuencia de ingestión, integración, ACV, tablas, figuras y documentos, y mantener la etiqueta provisional hasta completar las validaciones cruzadas.",
            "Priorizar, después de M3, una evaluación cuantitativa acotada del MCF del almacenamiento líquido y de los supuestos que controlan las etapas dominantes. Documentar como limitaciones los supuestos que no puedan resolverse con evidencia suficiente.",
            "Cotejar las referencias bibliográficas pendientes contra sus fuentes primarias y revisar institucionalmente la portada, los preliminares y la presentación final antes de generar el PDF de entrega.",
        ],
    )

    add_chapter(document, "9. Referencias")
    add_text(
        document,
        [
            "La bibliografía se consolida desde el registro integral versionado. La normalización editorial y la verificación de las entradas marcadas como pendientes deberán completarse antes de la versión final.",
        ],
    )
    for entry in references:
        paragraph = document.add_paragraph(entry["reference"], style="Normal")
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.left_indent = Inches(0.25)

    add_chapter(
        document,
        "Apéndice A. Trazabilidad entre objetivos y evidencia provisional",
    )
    traceability = pd.DataFrame(
        [
            [
                "Objetivo general",
                "Metodología ACV, comparación bajo unidad funcional común y evaluación por etapa",
                "Impactos por etapa, totales y comparación provisional M1–M2",
                "Conclusión provisional 5",
            ],
            [
                "Objetivo específico 1",
                "Caracterización, flujos, balances y construcción del inventario",
                "Caracterización, flujos y emisiones provisionales",
                "Conclusión provisional 1",
            ],
            [
                "Objetivo específico 2",
                "Caracterización EF 3.1 y agregación por etapa y escenario",
                "Cambio climático y eutrofización por etapa, escenario y unidad funcional",
                "Conclusiones provisionales 2 a 4",
            ],
        ],
        columns=["Objetivo", "Respuesta metodológica", "Evidencia integrada", "Conclusión relacionada"],
    )
    table_11 = add_dataframe(
        document,
        profile,
        counters,
        "Trazabilidad entre objetivos y evidencia provisional.",
        traceability,
        decimals=2,
    )
    add_text(
        document,
        [
            f"La Tabla {table_11} muestra cómo la metodología, los resultados y las conclusiones provisionales responden a los objetivos invariantes del estudio.",
        ],
    )

    finalize_document_format(document, profile)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    document.save(OUT_DOCX)
    return counters.table, counters.figure, counters.equation, len(references)


def all_document_text(document: Document) -> str:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + cells)


def validate_document(
    master_hash_before: str,
    master_hash_after: str,
    expected_tables: int,
    expected_figures: int,
    expected_equations: int,
    expected_references: int,
) -> None:
    with zipfile.ZipFile(OUT_DOCX) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise RuntimeError(f"El DOCX contiene un miembro ZIP corrupto: {bad_member}")

    document = Document(str(OUT_DOCX))
    text = all_document_text(document)
    checks: list[tuple[str, bool]] = [
        ("El DOCX abre como paquete válido", True),
        ("El objetivo general se conserva literalmente", OBJECTIVE_GENERAL in text),
        ("El primer objetivo específico se conserva literalmente", OBJECTIVE_SPECIFIC_1 in text),
        ("El segundo objetivo específico se conserva literalmente", OBJECTIVE_SPECIFIC_2 in text),
        ("La etiqueta PROVISIONAL M1–M2 es visible", PROVISIONAL_LABEL in text),
        ("M3 se identifica como pendiente", "M3" in text and "pendiente" in text.lower()),
        ("La jerarquía académica prevista está completa", all(title in text for title in EXPECTED_HEADINGS)),
        ("No hay marcadores accidentales", not re.search(r"\{\{|\}\}|\bTODO\b|\bTBD\b|Lorem ipsum|\[PENDIENTE\]", text)),
        ("No hay rutas internas visibles", not re.search(r"(?:processed|outputs|scripts|MASTER_escrito)[/\\]|\.csv\b", text, re.IGNORECASE)),
        (
            "No hay etiquetas técnicas internas prohibidas",
            not any(
                token.lower() in text.lower()
                for token in [
                    "dry_lot",
                    "uncovered_anaerobic_lagoon",
                    "composting_invessel",
                    "modelo_calculo",
                    "sistema_manejo_ipcc",
                    "n_ex_pct",
                    "n_ex_fraction",
                    "masa_total_kg_eq",
                    "hardcodeado",
                    "auditado",
                ]
            ),
        ),
        ("No hay errores visibles de codificación", not any(marker in text for marker in ["AnÃ", "metodologÃ", "estiÃ", "nitrÃ", "�"])),
        ("Las unidades anuales conservan la tilde", not re.search(r"/(?:ano|aNo)\b", text)),
        ("No hay etapas con decimales", not re.search(r"\b[AB][1-4][,.]0+\b", text)),
        ("No hay delimitadores visibles de ecuaciones", "\\[" not in text and "\\]" not in text and "$$" not in text),
        ("El MASTER conserva su hash registrado", master_hash_before == master_hash_after == REGISTERED_REFERENCE_SHA256),
        ("La salida está fuera del directorio protegido", MASTER.parent not in OUT_DOCX.parents),
        ("Las fuentes reproducibles del diagrama de fronteras existen", SYSTEM_BOUNDARY_PNG.exists() and SYSTEM_BOUNDARY_SVG.exists()),
        ("Las figuras insertadas coinciden con las previstas", len(document.inline_shapes) == expected_figures),
        ("Las tablas insertadas coinciden con las previstas", len(document.tables) == expected_tables),
        ("El Escenario A se identifica como operación habitual", "Escenario A" in text and "operación habitual" in text),
        ("El Escenario B se identifica como materializado temporalmente", "Escenario B" in text and "materializarlo temporalmente" in text),
        ("La tanqueta y el cañón tienen funciones distintas", "La tanqueta almacena" in text and "el cañón aplica" in text),
        ("Se distingue el intervalo previo al muestreo", all(term in text for term in ["jueves por la tarde", "lunes por la mañana", "entradas"])),
        ("A1 se describe sin precisión falsa", "21 días" in text and "tres a cuatro semanas" in text),
        ("A2 se describe como operación regular posterior a A1", "13 semanas" in text and "operación regular" in text and "después de A1" in text),
        ("No se atribuye una muestra de lombricompost terminado", "no muestreó lombricompost terminado" in text),
        ("El MCF se identifica como proxy no medido a tres días", "MCF de 38 %" in text and "proxy IPCC" in text and "No corresponde a una medición específica" in text),
        ("No se presenta 3,5 días como parámetro canónico", "3,5 días" not in text and "3.5 días" not in text),
    ]

    labels = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if re.fullmatch(r"(?:Tabla|Figura) \d+", paragraph.text.strip())
    ]
    checks.append(("No hay rótulos duplicados", len(labels) == len(set(labels))))
    checks.append(("La numeración global de tablas es continua", sorted(int(x.split()[1]) for x in labels if x.startswith("Tabla")) == list(range(1, expected_tables + 1))))
    checks.append(("La numeración global de figuras es continua", sorted(int(x.split()[1]) for x in labels if x.startswith("Figura")) == list(range(1, expected_figures + 1))))
    equation_numbers = [
        int(match.group(1))
        for paragraph in document.paragraphs
        if (match := re.search(r"\s+\((\d+)\)$", paragraph.text.strip()))
    ]
    checks.append(("La numeración global de ecuaciones es continua", equation_numbers == list(range(1, expected_equations + 1))))
    table_headers = [cell.text.strip() for table in document.tables for cell in table.rows[0].cells]
    checks.append(("Las tablas no duplican columnas de etapa", not ({"Etapa", "Nombre de etapa"} <= set(table_headers))))

    paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    clause_text = re.split(r"(?<=[.!?])\s+", paragraph_text)
    table_rows = [" ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    clauses_without_b = [
        re.sub(r"\bB[12]:.*$", "", clause, flags=re.IGNORECASE)
        for clause in clause_text + table_rows
    ]
    scenario_a_mislabel = any(
        re.search(r"\bA[1-4]:", clause)
        and re.search(r"\bpur[ií]n(?:es)?\b", clause, flags=re.IGNORECASE)
        for clause in clauses_without_b
    )
    checks.append(("No se asocian purines con etapas del Escenario A", not scenario_a_mislabel))

    failures = [name for name, passed in checks if not passed]
    status = "PASS" if not failures else "FAIL"
    lines = [
        "# Validación del documento integral provisional",
        "",
        f"- Estado general: **{status}**.",
        f"- Documento: `{OUT_DOCX.name}`.",
        f"- Tablas con numeración global: {expected_tables}.",
        f"- Figuras con numeración global: {expected_figures}.",
        f"- Ecuaciones con numeración global: {expected_equations}.",
        f"- Referencias integradas desde el registro bibliográfico: {expected_references}.",
        f"- SHA-256 del MASTER antes: `{master_hash_before}`.",
        f"- SHA-256 del MASTER después: `{master_hash_after}`.",
        "",
        "## Comprobaciones",
        "",
    ]
    lines.extend(f"- {'PASS' if passed else 'FAIL'} — {name}." for name, passed in checks)
    lines.extend(
        [
            "",
            "## Alcance de la validación",
            "",
            "- La validación comprueba estructura, objetivos, identificación provisional, integridad del paquete, numeración editorial, rutas visibles, figuras, tablas e integridad del MASTER.",
            "- La revisión visual y la validación científica supervisora permanecen separadas de estas comprobaciones programáticas.",
            "- La bibliografía conserva estados de verificación en el registro integral; las entradas pendientes requieren cotejo antes del cierre final.",
        ]
    )
    OUT_VALIDATION.write_text("\n".join(lines) + "\n", encoding="utf-8")
    if failures:
        raise RuntimeError(f"Falló la validación del documento integral: {failures}")


def main() -> None:
    validate_inputs()
    master_hash_before = sha256_file(MASTER)
    table_count, figure_count, equation_count, reference_count = build_document()
    master_hash_after = assert_reference_docx_intact(MASTER, master_hash_before)
    validate_document(
        master_hash_before,
        master_hash_after,
        table_count,
        figure_count,
        equation_count,
        reference_count,
    )
    print(f"Documento integral generado: {OUT_DOCX.relative_to(ROOT)}")
    print(f"Validación generada: {OUT_VALIDATION.relative_to(ROOT)}")
    print(f"MASTER sin cambios: {master_hash_after}")


if __name__ == "__main__":
    main()
