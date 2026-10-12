from __future__ import annotations

from docx import Document
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import parse_xml
from docx.shared import Inches

import latex2mathml.converter
import mathml2omml

MATH_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def latex_to_omml(latex_expression: str):
    """Convierte LaTeX canónico a un objeto matemático nativo de Word."""
    mathml = latex2mathml.converter.convert(latex_expression)
    omml = mathml2omml.convert(mathml).replace(
        "<m:oMath>", f'<m:oMath xmlns:m="{MATH_NAMESPACE}">', 1
    )
    return parse_xml(omml)


def add_word_equation(
    document: Document,
    latex_expression: str,
    *,
    number: int | None = None,
):
    """Inserta matemática OMML editable, centrada y con numeración robusta."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(3.0), WD_TAB_ALIGNMENT.CENTER
    )
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.15), WD_TAB_ALIGNMENT.RIGHT
    )
    paragraph.add_run("\t")
    paragraph._p.append(latex_to_omml(latex_expression))
    if number is not None:
        paragraph.add_run(f"\t({number})")
    return paragraph


def add_inline_math(paragraph: Paragraph, latex_expression: str):
    """Inserta un símbolo OMML editable dentro de un párrafo de prosa."""
    omml = latex_to_omml(latex_expression)
    paragraph._p.append(omml)
    return omml


def add_symbol_definitions(
    document: Document,
    definitions: list[tuple[str, str]],
    *,
    style: str = "Normal",
):
    """Añade un párrafo «Donde:» con símbolos OMML y definiciones académicas."""
    if not definitions:
        return None
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run("Donde: ").bold = True
    for index, (symbol_latex, explanation) in enumerate(definitions):
        if index:
            paragraph.add_run("; ")
        add_inline_math(paragraph, symbol_latex)
        paragraph.add_run(f" {explanation.rstrip('.;')}")
    paragraph.add_run(".")
    return paragraph
