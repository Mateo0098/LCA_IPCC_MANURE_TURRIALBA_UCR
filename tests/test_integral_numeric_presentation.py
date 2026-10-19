from __future__ import annotations

import re
import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "outputs" / "documentos_tfg" / "TFG_ACV_Estiercol_INTEGRAL_PROVISIONAL_M1_M2.docx"
PARTIAL_DOCX = [
    ROOT / "outputs" / "documentos_tfg" / "metodologia_desarrollada_tfg.docx",
    ROOT / "outputs" / "documentos_tfg" / "resultados_desarrollados_tfg.docx",
    ROOT / "outputs" / "documentos_tfg" / "conclusiones_desarrolladas_tfg.docx",
]


def visible_text(path: Path) -> str:
    document = Document(path)
    return "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    )


class IntegralNumericPresentationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.document = Document(DOCX)
        cls.text = visible_text(DOCX)
        cls.partial_text = "\n".join(visible_text(path) for path in PARTIAL_DOCX)

    def test_small_nonzero_results_remain_visible(self) -> None:
        self.assertIn("5,955 × 10⁻⁴", self.text)
        self.assertRegex(self.text, r"[+−]\d,\d{2} × 10⁻[⁰¹²³⁴⁵⁶⁷⁸⁹]+")

    def test_experimental_dispersion_is_explicit_and_not_overclaimed(self) -> None:
        self.assertIn("media ± desviación estándar entre campañas", self.text)
        self.assertIn("no constituye una incertidumbre expandida", self.text)

    def test_simapro_table_uses_observed_precision_and_scientific_differences(self) -> None:
        table = next(
            table
            for table in self.document.tables
            if table.rows[0].cells[0].text == "Escenario"
            and "SimaPro" in [cell.text for cell in table.rows[0].cells]
        )
        body = "\n".join(cell.text for row in table.rows[1:] for cell in row.cells)
        self.assertIn("3\u00a0167,9045", body)
        self.assertIn("−1,99 × 10⁻⁵", body)
        self.assertNotRegex(body, r"\d+[,.]\d{8,}")

    def test_no_negative_zero_or_decimal_stage_codes(self) -> None:
        self.assertNotRegex(self.text, r"(?<!\d)[−-]0(?:,0+)?(?:\s|$)")
        self.assertNotRegex(self.text, r"\b[AB][1-4][,.]0+\b")

    def test_exact_dates_and_codes_are_preserved(self) -> None:
        for value in ("2025", "2026", "A1", "B2", "M1", "M2"):
            self.assertIn(value, self.text)

    def test_derived_percentages_are_consistent_across_narrative_sections(self) -> None:
        for obsolete in ("53,93 %", "81,99 %", "70,63 %", "93,14 %", "62,92 %", "47,44 %"):
            self.assertNotIn(obsolete, self.text)
        for expected in ("53,9 %", "82 %", "70,6 %", "93,1 %", "62,9 %", "47,4 %"):
            self.assertIn(expected, self.text)

    def test_partial_documents_do_not_restore_known_excess_precision(self) -> None:
        for obsolete in (
            "0,333106919016",
            "3\u00a0170,113530",
            "0,120634221",
            "182,50 L/año",
            "26\u00a0278,73 kg/año",
        ):
            self.assertNotIn(obsolete, self.partial_text)
        self.assertIn("182,5 L/año", self.partial_text)


if __name__ == "__main__":
    unittest.main()
