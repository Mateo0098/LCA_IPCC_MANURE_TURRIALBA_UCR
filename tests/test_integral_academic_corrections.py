from __future__ import annotations

import sys
import unittest
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import generate_integral_tfg_docx as integral  # noqa: E402


class IntegralAcademicCorrectionsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.corrected_text = "\n".join(integral.MASTER_PARAGRAPH_OVERRIDES.values())

    def test_reported_regressions_are_absent_from_overrides(self) -> None:
        folded = self.corrected_text.casefold()
        for phrase in integral.ACADEMIC_REGRESSION_PHRASES:
            with self.subTest(phrase=phrase):
                self.assertNotIn(phrase.casefold(), folded)

    def test_corrected_narrative_citations_are_present(self) -> None:
        for citation in (
            "Vargas Sarmiento (2023)",
            "Rojas y Elizondo (2020)",
            "Komakech et al. (2015)",
            "International Organization for Standardization, 2006",
        ):
            with self.subTest(citation=citation):
                self.assertIn(citation, self.corrected_text)

    def test_climate_values_are_preserved_with_defined_meaning(self) -> None:
        paragraph = integral.MASTER_PARAGRAPH_OVERRIDES[100]
        self.assertIn("precipitación total anual promedio de 2.375,6 mm", paragraph)
        self.assertIn("promedios mensuales de temperatura máxima entre 28,2 °C y 28,5 °C", paragraph)

    def test_jjagwe_warning_preserves_approved_benchmark(self) -> None:
        paragraph = integral.MASTER_PARAGRAPH_OVERRIDES[86]
        self.assertIn("3,943 × 10⁻⁵ g N₂O/kg", paragraph)
        self.assertIn("39,43 mg N₂O/kg", paragraph)
        self.assertIn("sin tratarla como una validación del modelo", paragraph)

    def test_svg_hash_salt_is_stable(self) -> None:
        self.assertEqual(
            integral.SVG_HASHSALT,
            "lca-ipcc-manure-turrialba-ucr",
        )

    def test_page_numbering_is_continuous_and_hidden_on_cover(self) -> None:
        document = Document()
        document.add_section(WD_SECTION.NEW_PAGE)
        integral.configure_page_numbering(document)

        self.assertTrue(document.sections[0].different_first_page_header_footer)
        self.assertIn(
            integral.PROVISIONAL_LABEL,
            "\n".join(
                paragraph.text
                for paragraph in document.sections[0].first_page_header.paragraphs
            ),
        )
        first_page_instructions = [
            (node.text or "").strip()
            for paragraph in document.sections[0].first_page_footer.paragraphs
            for node in paragraph._p.iter(qn("w:instrText"))
        ]
        self.assertNotIn("PAGE", first_page_instructions)
        footer_instructions = [
            (node.text or "").strip()
            for paragraph in document.sections[0].footer.paragraphs
            for node in paragraph._p.iter(qn("w:instrText"))
        ]
        self.assertIn("PAGE", footer_instructions)
        self.assertTrue(document.sections[1].footer.is_linked_to_previous)

    def test_table_of_contents_is_native_navigable_and_automatically_updated(self) -> None:
        document = Document()
        marker = integral.add_preliminaries(document)

        instructions = [
            node.text or ""
            for paragraph in document.paragraphs
            for node in paragraph._p.iter(qn("w:instrText"))
            if "TOC" in (node.text or "")
        ]
        self.assertEqual(len(instructions), 1)
        self.assertIn('TOC \\o "1-3" \\h \\z \\u', instructions[0])
        update_fields = document.settings.element.find(qn("w:updateFields"))
        self.assertIsNotNone(update_fields)
        self.assertEqual(update_fields.get(qn("w:val")), "true")

        toc_titles = [
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.strip() == "Contenido"
        ]
        self.assertEqual(len(toc_titles), 1)
        self.assertEqual(toc_titles[0].style.name, integral.TOC_STYLE_NAME)
        outline_level = toc_titles[0].style.element.get_or_add_pPr().find(
            qn("w:outlineLvl")
        )
        self.assertIsNotNone(outline_level)
        self.assertEqual(outline_level.get(qn("w:val")), "9")
        preliminary_text = [paragraph.text for paragraph in document.paragraphs]
        self.assertFalse(
            any(heading in preliminary_text for heading in integral.EXPECTED_HEADINGS)
        )
        self.assertEqual(marker.text, "__LISTA_SIGLAS__")

    def test_nitrogen_contrast_table_supports_positive_negative_and_zero(self) -> None:
        context = {
            "A1": {"reference": 8.0, "model": 10.0, "difference": 2.0, "relative": 25.0},
            "A3": {"reference": 0.012, "model": 0.010, "difference": -0.002, "relative": -16.666},
            "B1": {"reference": 0.02, "model": 0.02, "difference": 0.0, "relative": 0.0},
        }
        table = integral.nitrogen_contrast_table(context)

        self.assertEqual(table.iloc[0]["Diferencia absoluta (modelo − referencia)"], "+2,0 kg N/año")
        self.assertEqual(table.iloc[1]["Diferencia absoluta (modelo − referencia)"], "−0,0020 puntos porcentuales")
        self.assertEqual(table.iloc[2]["Diferencia absoluta (modelo − referencia)"], "0,0000 puntos porcentuales")
        self.assertEqual(table.iloc[0]["Diferencia relativa"], "+25,0 %")
        self.assertEqual(table.iloc[1]["Diferencia relativa"], "−16,7 %")
        self.assertEqual(table.iloc[2]["Diferencia relativa"], "0,0 %")

    def test_nitrogen_narrative_tracks_positive_negative_and_zero(self) -> None:
        context = {
            "A1": {"reference": 8.0, "model": 10.0, "difference": 2.0, "relative": 25.0},
            "A3": {"reference": 0.012, "model": 0.010, "difference": -0.002, "relative": -16.666},
            "B1": {"reference": 0.02, "model": 0.02, "difference": 0.0, "relative": 0.0},
        }
        results = "\n".join(integral.nitrogen_results_paragraphs(context))
        discussion = "\n".join(integral.nitrogen_discussion_paragraphs(context))
        conclusion = integral.nitrogen_conclusion_text(context)

        self.assertIn("conservó más N residual", results)
        self.assertIn("concentración modelada fue menor", results)
        self.assertIn("no se distinguió de la medición M2", results)
        self.assertIn("superó a la referencia derivada", discussion)
        self.assertIn("16,7 % menor", discussion)
        self.assertIn("no mostró una diferencia visible", discussion)
        self.assertIn(
            "superó la referencia en A1, fue inferior a la referencia en aguas verdes y no mostró diferencia visible respecto de la referencia en purines",
            conclusion,
        )

    def test_rounded_zero_has_no_false_positive_or_negative_sign(self) -> None:
        self.assertEqual(
            integral.nitrogen_direction_wording(
                0.004,
                2,
                positive="positivo",
                negative="negativo",
                zero="cero",
            ),
            "cero",
        )
        self.assertEqual(
            integral.format_spanish_number(-0.004, 2, explicit_sign=True),
            "0,00",
        )


if __name__ == "__main__":
    unittest.main()
